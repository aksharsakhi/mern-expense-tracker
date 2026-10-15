import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Bespoke Color Palette
    BG_COLOR = RGBColor(15, 23, 42)          # Deep Slate #0f172a
    CARD_BG = RGBColor(30, 41, 59)           # Card Slate #1e293b
    CARD_BORDER = RGBColor(51, 65, 85)       # Subtle border #334155
    ACCENT_INDIGO = RGBColor(99, 102, 241)   # Primary Indigo #6366f1
    ACCENT_EMERALD = RGBColor(16, 185, 129)  # Income Emerald #10b981
    ACCENT_ROSE = RGBColor(244, 63, 94)      # Expense Rose #f43f5e
    ACCENT_AMBER = RGBColor(245, 158, 11)    # Amber #f59e0b
    TEXT_MAIN = RGBColor(248, 250, 252)      # #f8fafc
    TEXT_MUTED = RGBColor(148, 163, 184)     # #94a3b8

    def apply_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category="FULL STACK FRAMEWORKS — MINI PROJECT"):
        # Header category badge
        badge = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.4))
        tf = badge.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = category.upper()
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = ACCENT_INDIGO

        # Main slide title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.7), Inches(0.8))
        tf2 = title_box.text_frame
        tf2.word_wrap = True
        p2 = tf2.paragraphs[0]
        p2.text = title_text
        p2.font.size = Pt(26)
        p2.font.bold = True
        p2.font.color.rgb = TEXT_MAIN

    # ==========================================================
    # SLIDE 1: TITLE SLIDE
    # ==========================================================
    s1 = prs.slides.add_slide(blank_layout)
    apply_background(s1)

    # Accent decorative banner
    bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.8), Inches(0.12), Inches(3.6))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT_INDIGO
    bar.line.fill.background()

    # Title & Subtitle box
    tbox = s1.shapes.add_textbox(Inches(1.2), Inches(1.6), Inches(11.0), Inches(3.8))
    tf = tbox.text_frame
    tf.word_wrap = True

    p_badge = tf.paragraphs[0]
    p_badge.text = "FULL STACK FRAMEWORKS MINI-PROJECT EVALUATION"
    p_badge.font.size = Pt(13)
    p_badge.font.bold = True
    p_badge.font.color.rgb = ACCENT_EMERALD

    p_title = tf.add_paragraph()
    p_title.text = "SpendWise: Modern MERN Expense Tracker"
    p_title.font.size = Pt(36)
    p_title.font.bold = True
    p_title.font.color.rgb = TEXT_MAIN
    p_title.space_before = Pt(10)

    p_sub = tf.add_paragraph()
    p_sub.text = "Daily Cashflow Management, Monthly Aggregations & Visual Analytics"
    p_sub.font.size = Pt(18)
    p_sub.font.color.rgb = TEXT_MUTED
    p_sub.space_before = Pt(8)

    # Metadata card
    meta_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(5.4), Inches(10.8), Inches(1.2))
    meta_card.fill.solid()
    meta_card.fill.fore_color.rgb = CARD_BG
    meta_card.line.color.rgb = CARD_BORDER

    mtf = meta_card.text_frame
    mtf.word_wrap = True
    mp = mtf.paragraphs[0]
    mp.text = "• Candidate: Akshar Sakhi      • Evaluation Date: 16 October 2026      • Stack: MongoDB, Express, React, Node.js"
    mp.font.size = Pt(13)
    mp.font.bold = True
    mp.font.color.rgb = TEXT_MAIN
    mp.alignment = PP_ALIGN.CENTER

    # ==========================================================
    # SLIDE 2: PROBLEM STATEMENT & CORE OBJECTIVES
    # ==========================================================
    s2 = prs.slides.add_slide(blank_layout)
    apply_background(s2)
    add_header(s2, "Problem Statement & Project Objectives")

    col_w = Inches(3.64)
    h_card = Inches(4.8)
    top_pos = Inches(1.8)

    cards_data = [
      (
        "1. The Challenge",
        ACCENT_ROSE,
        [
          "Personal budgeting is often cumbersome and dispersed across spreadsheets or paper receipts.",
          "Lack of real-time visibility into category-specific spending leaks (e.g. food, utilities, housing).",
          "Absence of instant monthly cashflow aggregations to determine true savings rates."
        ]
      ),
      (
        "2. Our Solution",
        ACCENT_INDIGO,
        [
          "SpendWise: A centralized, high-speed MERN web platform.",
          "Seamless daily recording of both income and expenses with payment method and notes.",
          "Interactive monthly date pickers with live multi-dimensional analytics and instant charts."
        ]
      ),
      (
        "3. Core Objectives",
        ACCENT_EMERALD,
        [
          "Full RESTful CRUD operations with server-side validation.",
          "High-performance MongoDB aggregation pipelines for instant monthly totals.",
          "Responsive, glassmorphic React dashboard with dark/light theme support."
        ]
      )
    ]

    for i, (ctitle, ccolor, points) in enumerate(cards_data):
      left = Inches(0.8 + i * 4.0)
      card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top_pos, col_w, h_card)
      card.fill.solid()
      card.fill.fore_color.rgb = CARD_BG
      card.line.color.rgb = CARD_BORDER

      ctf = card.text_frame
      ctf.word_wrap = True
      cp = ctf.paragraphs[0]
      cp.text = ctitle
      cp.font.size = Pt(16)
      cp.font.bold = True
      cp.font.color.rgb = ccolor
      cp.space_before = Pt(8)

      for pt in points:
        bp = ctf.add_paragraph()
        bp.text = f"•  {pt}"
        bp.font.size = Pt(12)
        bp.font.color.rgb = TEXT_MUTED
        bp.space_before = Pt(12)

    # ==========================================================
    # SLIDE 3: SYSTEM ARCHITECTURE & TECH STACK
    # ==========================================================
    s3 = prs.slides.add_slide(blank_layout)
    apply_background(s3)
    add_header(s3, "Full-Stack Architecture & Technology Stack")

    tiers = [
      ("Frontend Tier", ACCENT_INDIGO, [
        "React 18 & Vite: Ultra-fast component rendering and HMR.",
        "Chart.js & React-Chartjs-2: Interactive category donut and daily cashflow bar graphs.",
        "Lucide React: Crisp vector iconography for all financial categories.",
        "Design System: Modern CSS variables, glassmorphism, fluid micro-interactions, dark/light themes."
      ]),
      ("Backend REST Tier", ACCENT_AMBER, [
        "Node.js & Express.js: Lightweight, modular HTTP server architecture.",
        "REST API Contract: Clean separation of concerns with controllers, models, and routes.",
        "CORS & Morgan: Cross-origin security and structured request logging.",
        "Centralized Error Handler: Uniform JSON error reporting across all endpoints."
      ]),
      ("Database & Aggregation Tier", ACCENT_EMERALD, [
        "MongoDB 8.x & Mongoose: Document-based database with strict schema validation.",
        "Multi-Faceted Aggregation Pipelines ($facet, $group, $sort): Calculates totals in a single DB round-trip.",
        "Performance Indexes: Compound indexes on (date, type) and category for instant query retrieval."
      ])
    ]

    for i, (ttitle, tcolor, bullets) in enumerate(tiers):
      top = Inches(1.8 + i * 1.7)
      tcard = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top, Inches(11.73), Inches(1.45))
      tcard.fill.solid()
      tcard.fill.fore_color.rgb = CARD_BG
      tcard.line.color.rgb = CARD_BORDER

      ttf = tcard.text_frame
      ttf.word_wrap = True
      tp = ttf.paragraphs[0]
      tp.text = ttitle
      tp.font.size = Pt(15)
      tp.font.bold = True
      tp.font.color.rgb = tcolor

      for b in bullets:
        bp = ttf.add_paragraph()
        bp.text = f"• {b}"
        bp.font.size = Pt(11)
        bp.font.color.rgb = TEXT_MUTED
        bp.space_before = Pt(3)

    # ==========================================================
    # SLIDE 4: REST API DESIGN & DATA CONTRACTS
    # ==========================================================
    s4 = prs.slides.add_slide(blank_layout)
    apply_background(s4)
    add_header(s4, "RESTful API Specification & JSON Data Contract")

    # Left: Endpoints table card
    api_card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(6.8), Inches(4.8))
    api_card.fill.solid()
    api_card.fill.fore_color.rgb = CARD_BG
    api_card.line.color.rgb = CARD_BORDER

    atf = api_card.text_frame
    atf.word_wrap = True
    ap = atf.paragraphs[0]
    ap.text = "API Endpoint Contracts"
    ap.font.size = Pt(16)
    ap.font.bold = True
    ap.font.color.rgb = ACCENT_INDIGO

    endpoints = [
      ("GET /api/transactions?month=2026-10", "Fetches filtered list by month, type, category, and keyword search"),
      ("GET /api/transactions/summary", "MongoDB aggregation computing totals, savings rate, and category breakdowns"),
      ("POST /api/transactions", "Creates transaction with validation (amount, type, category, date, method)"),
      ("PUT /api/transactions/:id", "Updates transaction fields with Mongoose validator checks"),
      ("DELETE /api/transactions/:id", "Permanently deletes transaction record by ObjectId"),
      ("POST /api/transactions/seed", "Instantly populates realistic evaluation demo dataset")
    ]

    for ep, desc in endpoints:
      p1 = atf.add_paragraph()
      p1.text = ep
      p1.font.size = Pt(11)
      p1.font.bold = True
      p1.font.color.rgb = ACCENT_EMERALD
      p1.space_before = Pt(6)

      p2 = atf.add_paragraph()
      p2.text = f"  └ {desc}"
      p2.font.size = Pt(10)
      p2.font.color.rgb = TEXT_MUTED

    # Right: JSON Sample card
    json_card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.8), Inches(1.8), Inches(4.7), Inches(4.8))
    json_card.fill.solid()
    json_card.fill.fore_color.rgb = CARD_BG
    json_card.line.color.rgb = CARD_BORDER

    jtf = json_card.text_frame
    jtf.word_wrap = True
    jp = jtf.paragraphs[0]
    jp.text = "Mongoose Transaction Schema"
    jp.font.size = Pt(16)
    jp.font.bold = True
    jp.font.color.rgb = ACCENT_AMBER

    json_code = (
      "{\n"
      '  "_id": "670ecb3523f4...",\n'
      '  "amount": 18500,\n'
      '  "type": "expense",\n'
      '  "category": "Housing",\n'
      '  "date": "2026-10-02T10:30:00Z",\n'
      '  "paymentMethod": "Net Banking",\n'
      '  "note": "Apartment monthly rent",\n'
      '  "createdAt": "2026-10-02T...",\n'
      '  "updatedAt": "2026-10-02T..."\n'
      "}\n\n"
      "Validation Rules:\n"
      "• amount > 0 (numeric min: 0.01)\n"
      "• type: ['income', 'expense']\n"
      "• category: string (required)\n"
      "• paymentMethod: enum (UPI, Net Banking, Card, Cash)"
    )

    jp2 = jtf.add_paragraph()
    jp2.text = json_code
    jp2.font.size = Pt(10.5)
    jp2.font.color.rgb = TEXT_MAIN
    jp2.space_before = Pt(8)

    # ==========================================================
    # SLIDE 5: MONGODB AGGREGATION HIGHLIGHTS
    # ==========================================================
    s5 = prs.slides.add_slide(blank_layout)
    apply_background(s5)
    add_header(s5, "MongoDB Aggregation Pipeline Architecture")

    agg_cards = [
      ("Pipeline Stage 1: $match", ACCENT_INDIGO, [
        "Constructs precise UTC date boundaries for the selected month (e.g., 2026-10-01 to 2026-11-01).",
        "Filters the indexed date field in logarithmic time before branching into parallel facets."
      ]),
      ("Pipeline Stage 2: Multi-Faceted $facet", ACCENT_EMERALD, [
        "totals: Computes overall income sum, expense sum, and transaction counts.",
        "expensesByCategory: Groups expenses by category and sorts descending by total expenditure.",
        "dailyTotals: Groups amounts by day and type for visual time-series velocity graphs.",
        "paymentMethods: Analyzes distribution across UPI, Cards, and Net Banking."
      ]),
      ("Pipeline Stage 3: Computed Metrics", ACCENT_AMBER, [
        "Net Balance: Evaluated in real-time as (totalIncome - totalExpense).",
        "Savings Rate: Evaluated as ((income - expense) / income) * 100.",
        "Category Share: Real-time percentage calculation per category."
      ])
    ]

    for i, (title, color, bullets) in enumerate(agg_cards):
      left = Inches(0.8 + i * 4.0)
      acard = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.8), Inches(3.64), Inches(4.8))
      acard.fill.solid()
      acard.fill.fore_color.rgb = CARD_BG
      acard.line.color.rgb = CARD_BORDER

      atf = acard.text_frame
      atf.word_wrap = True
      ap = atf.paragraphs[0]
      ap.text = title
      ap.font.size = Pt(15)
      ap.font.bold = True
      ap.font.color.rgb = color
      ap.space_before = Pt(8)

      for b in bullets:
        bp = atf.add_paragraph()
        bp.text = f"•  {b}"
        bp.font.size = Pt(12)
        bp.font.color.rgb = TEXT_MUTED
        bp.space_before = Pt(12)

    # ==========================================================
    # SLIDE 6: DEMO HIGHLIGHTS & CONCLUSION
    # ==========================================================
    s6 = prs.slides.add_slide(blank_layout)
    apply_background(s6)
    add_header(s6, "Evaluation Demonstration & Key Takeaways")

    demo_card = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.7), Inches(4.8))
    demo_card.fill.solid()
    demo_card.fill.fore_color.rgb = CARD_BG
    demo_card.line.color.rgb = CARD_BORDER

    dtf = demo_card.text_frame
    dtf.word_wrap = True
    dp = dtf.paragraphs[0]
    dp.text = "Live Demonstration Workflow"
    dp.font.size = Pt(16)
    dp.font.bold = True
    dp.font.color.rgb = ACCENT_INDIGO

    steps = [
      ("1. Dynamic Month Switching", "Observe instant recalculation of metrics when switching between October and September 2026."),
      ("2. Full CRUD Operations", "Add a new transaction via modal, edit amounts/notes on the fly, or delete with confirmation modal."),
      ("3. Multi-Faceted Filters", "Filter transactions live by search keyword, income/expense type, and category."),
      ("4. One-Click Demo Seeder", "Single-click 'Seed Data' button to refresh realistic evaluation data at any time."),
      ("5. CSV Export", "Direct export of month records to standard CSV file format.")
    ]

    for st, sd in steps:
      sp1 = dtf.add_paragraph()
      sp1.text = st
      sp1.font.size = Pt(11)
      sp1.font.bold = True
      sp1.font.color.rgb = ACCENT_EMERALD
      sp1.space_before = Pt(6)

      sp2 = dtf.add_paragraph()
      sp2.text = f"   └ {sd}"
      sp2.font.size = Pt(10)
      sp2.font.color.rgb = TEXT_MUTED

    # Right: Summary & Learning Outcomes
    conc_card = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8))
    conc_card.fill.solid()
    conc_card.fill.fore_color.rgb = CARD_BG
    conc_card.line.color.rgb = CARD_BORDER

    ctf2 = conc_card.text_frame
    ctf2.word_wrap = True
    cp2 = ctf2.paragraphs[0]
    cp2.text = "Key Learning Outcomes"
    cp2.font.size = Pt(16)
    cp2.font.bold = True
    cp2.font.color.rgb = ACCENT_AMBER

    learnings = [
      "Mastery of the MERN architecture with clean API contracts and data models.",
      "Hands-on expertise building advanced MongoDB aggregation pipelines ($facet, $group, $match).",
      "Building responsive, accessible React applications with interactive Chart.js visualizations.",
      "Adhering to professional Git development workflows with clear, atomic commit histories.",
      "Clean separation of concerns: reusable UI components, centralized API services, and robust server error handling."
    ]

    for l in learnings:
      lp = ctf2.add_paragraph()
      lp.text = f"✔  {l}"
      lp.font.size = Pt(11.5)
      lp.font.color.rgb = TEXT_MAIN
      lp.space_before = Pt(10)

    # Save presentation
    out_dir = os.path.join(os.path.dirname(__file__))
    out_path = os.path.join(out_dir, "Mini_Project_Evaluation_Expense_Tracker.pptx")
    prs.save(out_path)
    print(f"✅ Presentation successfully generated at: {out_path}")

if __name__ == "__main__":
    create_deck()
